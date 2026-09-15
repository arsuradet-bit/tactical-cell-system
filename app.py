from __future__ import annotations
import hmac
import os
import time
from io import BytesIO
from datetime import datetime, date, time as day_time
from pathlib import Path
import pandas as pd
import streamlit as st
from streamlit_folium import st_folium
import survey_database as db
from core import NETWORKS, network_label, paired_ids, parse_ids, prepare_cdr, prepare_gmon, read_table, match_cdr, display_time, coordinates
try:
    from core import prepare_camera
except ImportError:
    prepare_camera = None
from survey_maps import build_map

ROOT = Path(__file__).parent
st.set_page_config(page_title="วิเคราะห์พิกัด", page_icon=str(ROOT/"logo-transparent.png") if (ROOT/"logo-transparent.png").exists() else "📍", layout="wide", initial_sidebar_state="expanded")
for name in ("fonts.css", "style.css"):
    st.markdown('<style>' + (ROOT/name).read_text(encoding="utf-8") + '</style>', unsafe_allow_html=True)
DEMO = os.getenv("INTEL_DEMO") == "1"
if st.session_state.get("workspace_version") != 2:
    for old_key in list(st.session_state):
        if old_key != "authenticated":
            del st.session_state[old_key]
    st.session_state.workspace_version = 2


def logout():
    for key in list(st.session_state):
        del st.session_state[key]


def clear_search():
    for key in ("cell_input", "lac_input", "nbid_input"):
        st.session_state[key] = ""
    st.session_state.plmn_input = "ทุกเครือข่าย"
    st.session_state.update(rows=[], mode="ค้นหา", notice="", elapsed=None)
    st.session_state.revision += 1


def clear_cdr():
    st.session_state.cdr_generation += 1
    st.session_state.pop("cdr_errors", None)
    if st.session_state.mode == "CDR":
        st.session_state.update(rows=[], mode="ค้นหา", notice="", elapsed=None)
    st.session_state.revision += 1


try:
    password_expected = st.secrets.get("APP_PASSWORD", "")
except (FileNotFoundError, st.errors.StreamlitSecretNotFoundError):
    password_expected = ""
if not DEMO and not st.session_state.get("authenticated"):
    _, middle, _ = st.columns([1, 1.3, 1])
    with middle:
        if (ROOT/"logo-transparent.png").exists():
            st.image(str(ROOT/"logo-transparent.png"), width=110)
        st.title("วิเคราะห์พิกัด")
        st.caption("เข้าสู่พื้นที่วิเคราะห์ภายในองค์กร")
        with st.form("login"):
            password = st.text_input("รหัสผ่านองค์กร", type="password")
            login = st.form_submit_button("เข้าสู่ระบบ", type="primary", width="stretch")
        if login:
            if password_expected and hmac.compare_digest(password.encode(), str(password_expected).encode()):
                st.session_state.authenticated = True
                st.rerun()
            else:
                st.error("เข้าสู่ระบบไม่สำเร็จ กรุณาตรวจรหัสผ่าน")
    st.stop()


@st.cache_resource
def engine_resource():
    return db.connect(dict(st.secrets["database"]))


@st.cache_data(ttl=30)
def health(_engine):
    return db.health(_engine)


engine, ready = None, False
if not DEMO:
    try:
        engine = engine_resource()
        ready = health(engine)
    except Exception:
        pass
for key, default in {"rows": [], "mode": "ค้นหา", "revision": 0, "cdr_generation": 0, "notice": "", "elapsed": None}.items():
    if key not in st.session_state:
        st.session_state[key] = default

with st.sidebar:
    if (ROOT/"logo-transparent.png").exists():
        st.image(str(ROOT/"logo-transparent.png"), width=68)
    st.markdown('<div class="eyebrow">FIELD INTELLIGENCE</div><div class="brand">วิเคราะห์พิกัด</div>', unsafe_allow_html=True)
    st.caption("ระบบวิเคราะห์ภาคสนาม")
    workspace_menu = st.radio("เมนูหลัก", ["🔎 ค้นหา LAC / CELL / xNBID", "📥 เพิ่มข้อมูล G-Mon", "📞 วิเคราะห์ CDR", "🎥 กล้อง + CDR", "🗺️ ค้นหาพื้นที่"], key="workspace_menu", label_visibility="collapsed")
    st.divider()
    st.caption("● Cloud SQL พร้อมใช้งาน" if ready else "○ ตัวอย่างสังเคราะห์" if DEMO else "○ Cloud SQL ยังไม่พร้อม")
    st.button("ออกจากระบบ", on_click=logout, width="stretch")
active_panel = {"🔎 ค้นหา LAC / CELL / xNBID": "search_panel", "📥 เพิ่มข้อมูล G-Mon": "survey_upload", "📞 วิเคราะห์ CDR": "cdr_panel", "🎥 กล้อง + CDR": "camera_panel", "🗺️ ค้นหาพื้นที่": "area_panel"}[workspace_menu]
# Each workspace keeps its own result set; uploaded files stay session-only.
previous_panel = st.session_state.get("result_panel")
if previous_panel != active_panel:
    if previous_panel:
        st.session_state["result_" + previous_panel] = {k: st.session_state[k] for k in ("rows", "mode", "notice", "elapsed")}
    st.session_state.update(st.session_state.get("result_" + active_panel, dict(rows=[], mode="ค้นหา", notice="", elapsed=None)))
    st.session_state.result_panel = active_panel
    st.session_state.revision += 1
camera_focus_rows = None
camera_focus_key = "all"
# Keep upload widgets mounted when navigating, so session-only files survive.
hidden_panels = [name for name in ("search_panel", "survey_upload", "cdr_panel", "camera_panel", "area_panel") if name != active_panel]
st.markdown("<style>" + ",".join(".st-key-" + name for name in hidden_panels) + "{display:none}</style>", unsafe_allow_html=True)
st.markdown('<div class="workspace-heading"><div><div class="eyebrow">FIELD WORKSPACE</div><h2>' + workspace_menu + '</h2><p>ค้นหา เชื่อมโยง และตรวจสอบเหตุการณ์บนแผนที่</p></div><span class="workspace-badge">พื้นที่วิเคราะห์</span></div>', unsafe_allow_html=True)
control_column, map_column = st.columns([1.2, 2], gap="medium")
controls = control_column.container(key="workflow_controls")
map_area = map_column.container(key="map_workspace")
report_area = st.container(key="timeline_workspace")
if DEMO:
    st.info("ข้อมูลสังเคราะห์สำหรับทดสอบ · ไม่บันทึกฐานข้อมูล")
elif not ready:
    st.warning("ยังค้นหาหรือเพิ่มข้อมูลไม่ได้ กรุณาตรวจการเชื่อมต่อและติดตั้ง schema_all_points.sql")

# Upload is visible near the top, not hidden in an expander.
with controls.container(border=True, key="survey_upload"):
    st.markdown("### 📥 อัปโหลดข้อมูล G-Mon Pro")
    st.caption("ทุกคนเพิ่มข้อมูลได้ · เก็บทุกแถวรวมรายการซ้ำ · ไม่จำกัดจำนวนแถว")
    upload = st.file_uploader("เลือกไฟล์ G-Mon Pro", type=["csv", "txt"], key="gmon_upload")
    if upload:
        try:
            frame = read_table(upload.getvalue(), upload.name)
            records, issues = prepare_gmon(frame)
            st.caption(f"ข้อมูลทั้งหมด {len(frame):,} แถว · ผ่านการตรวจ {len(records):,} แถว · ต้องตรวจสอบ {len(issues):,} แถว")
            with st.expander("ตรวจข้อมูลก่อนบันทึก"):
                st.dataframe(frame.head(20), hide_index=True)
                if not issues.empty:
                    st.dataframe(issues, hide_index=True)
            if not issues.empty:
                st.warning("บันทึกต้นฉบับครบทุกแถว แต่แถวที่ไม่มีพิกัดใช้ได้จะแสดงเฉพาะในตารางผลค้นหา")
            if st.button("บันทึก G-Mon เข้าฐานข้อมูล", type="primary", disabled=not ready):
                with st.spinner("กำลังบันทึกทุกแถว…"):
                    count = db.save_raw(engine, frame)
                st.success(f"บันทึก {count:,} แถว รวมรายการซ้ำเรียบร้อย กดซ้ำจะเพิ่มข้อมูลอีกรอบ")
        except ValueError as exc:
            st.error(str(exc))
        except Exception:
            st.error("บันทึกไม่สำเร็จ ไม่มีการบันทึกบางส่วน กรุณาตรวจไฟล์และการเชื่อมต่อ")


def demo_rows():
    return [dict(observation_key=f"demo:{i}", plmn="52003", xci="350342102" if i < 3 else "117266",
                 lac="23072" if i < 3 else "55653", xnbid="116876", system="4", lat=7.01+i*.001,
                 lon=100.47+i*.001, rsrp=[-82,-98,-112,-88,-102][i], observed_at=datetime(2025,9,3,9,11,i)) for i in range(5)]


with controls.container(border=True, key="search_panel"):
    st.markdown("### 🔎 ค้นหาพิกัดภาคสนาม")
    with st.form("search"):
        a, b, c = st.columns(3)
        cells = a.text_area("CELL / XCI", placeholder="117266\n116876\n64918", height=140, key="cell_input")
        lacs = b.text_area("LAC / TAC", placeholder="55653\n55653\n06611", height=140, key="lac_input")
        nbids = c.text_area("xNBID", placeholder="116876\n117266", height=140, key="nbid_input")
        plmn = st.selectbox("เครือข่าย / PLMN", ["ทุกเครือข่าย"] + list(NETWORKS), format_func=network_label, key="plmn_input")
        st.caption("CELL และ LAC จับคู่ตามบรรทัด · เว้นทั้งช่องเพื่อค้นหาอีกช่องอย่างเดียว · xNBID ใส่ได้หลายบรรทัด")
        search_button = st.form_submit_button("ค้นหาพิกัด", type="primary", width="stretch", disabled=not (ready or DEMO))
    st.button("ล้างการค้นหา", on_click=clear_search, width="stretch", key="clear_search_button")
    if search_button:
        try:
            c, l, pairs = paired_ids(cells, lacs)
            n = parse_ids(nbids)
            if not any((c,l,n)):
                raise ValueError("กรุณากรอก CELL, LAC หรือ xNBID อย่างน้อยหนึ่งช่อง")
            begin = time.perf_counter()
            with st.spinner("กำลังค้นหาทุกจุดจากข้อมูลต้นฉบับ…"):
                if DEMO:
                    rows = [r for r in demo_rows() if (not pairs or (r['xci'],r['lac']) in pairs) and
                            (pairs or ((not c or r['xci'] in c) and (not l or r['lac'] in l))) and
                            (not n or r['xnbid'] in n) and (plmn == "ทุกเครือข่าย" or r['plmn'] == plmn)]
                else:
                    rows, _ = db.search(engine, c, l, n, [] if plmn == "ทุกเครือข่าย" else [plmn], pairs=pairs)
            st.session_state.update(rows=rows, mode="ค้นหา", elapsed=time.perf_counter()-begin,
                                    notice=f"พบ {len(rows):,} รายการ รวมข้อมูลซ้ำ")
            st.session_state.revision += 1
        except ValueError as exc:
            st.error(str(exc))
        except Exception:
            st.error("ค้นหาไม่สำเร็จ กรุณาตรวจการเชื่อมต่อแล้วลองใหม่")

with controls.container(key="cdr_panel"):
    st.caption("VOICE / DATA · ใช้ชั่วคราว")
    st.caption("จับคู่ทุกพิกัด G-Mon ด้วย LAC/CELL · ใช้พิกัด CDR เมื่อไม่พบ G-Mon · ไม่บันทึก CDR ลงฐานข้อมูล")
    a, b = st.columns(2)
    voice = a.file_uploader("ไฟล์ VOICE", type=["csv","txt","xlsx"], key=f"voice_{st.session_state.cdr_generation}")
    data = b.file_uploader("ไฟล์ DATA", type=["csv","txt","xlsx"], key=f"data_{st.session_state.cdr_generation}")
    with st.form("cdr_options"):
        cdr_network = st.selectbox("เครือข่ายของ CDR", ["ไม่ระบุ"]+list(NETWORKS), format_func=network_label)
        filter_time = st.checkbox("กรองช่วงเวลาเหตุการณ์")
        a,b = st.columns(2)
        start_day = a.date_input("วันที่เริ่ม", date.today())
        start_time = a.time_input("เวลาเริ่ม", day_time(0,0))
        end_day = b.date_input("วันที่สิ้นสุด", date.today())
        end_time = b.time_input("เวลาสิ้นสุด", day_time(23,59,59))
        analyze = st.form_submit_button("วิเคราะห์ VOICE / DATA", type="primary", disabled=not (ready or DEMO))
    st.button("ล้างไฟล์และผล CDR", on_click=clear_cdr)
    if analyze:
        try:
            if not voice and not data:
                raise ValueError("เลือกไฟล์ VOICE หรือ DATA อย่างน้อยหนึ่งไฟล์")
            batches, all_errors, offset = [], [], 0
            for kind, file in [("VOICE",voice),("DATA",data)]:
                if file:
                    raw = read_table(file.getvalue(), file.name, max_rows=None)
                    events, errors = prepare_cdr(raw, kind=kind, offset=offset)
                    offset += len(raw)
                    if not events.empty:
                        batches.append(events)
                    if not errors.empty:
                        errors["ประเภท"] = kind
                        all_errors.append(errors)
            st.session_state.cdr_errors = pd.concat(all_errors, ignore_index=True) if all_errors else pd.DataFrame()
            if not batches:
                raise ValueError("ไม่พบเหตุการณ์ที่อ่านวันเวลาและ LAC/CELL ได้")
            events = pd.concat(batches, ignore_index=True).sort_values(["event_at","event_id"])
            if filter_time:
                start, end = datetime.combine(start_day,start_time), datetime.combine(end_day,end_time)
                if start > end:
                    raise ValueError("เวลาเริ่มต้องไม่เกินเวลาสิ้นสุด")
                events = events[events.event_at.between(start,end)]
            if events.empty:
                raise ValueError("ไม่พบเหตุการณ์ในช่วงเวลาที่เลือก กรุณาตรวจช่วงเวลาหรือปิดตัวกรองเวลา")
            with st.spinner("กำลังจับคู่ทุกจุด G-Mon…"):
                candidates = demo_rows() if DEMO else db.lookup_pairs(engine,list(zip(events.xci,events.lac)),[] if cdr_network == "ไม่ระบุ" else [cdr_network])
                if DEMO and cdr_network != "ไม่ระบุ":
                    candidates = [r for r in candidates if r['plmn']==cdr_network]
                rows = match_cdr(events,candidates)
            st.session_state.update(rows=rows,mode="CDR",notice=f"วิเคราะห์ {len(events):,} เหตุการณ์ · จับคู่ได้ {len(rows):,} รายการ",elapsed=None)
            st.session_state.revision += 1
        except ValueError as exc:
            st.error(str(exc))
        except Exception as exc:
            # Surface a useful, non-secret diagnostic instead of hiding whether
            # the failure came from the file schema or the database connection.
            detail = str(exc).replace(str(st.secrets.get("database", {}).get("password", "")), "[ซ่อน]")
            st.error(f"วิเคราะห์ไม่สำเร็จ: {type(exc).__name__} — {detail or 'ไม่ทราบสาเหตุ'}")
    if not st.session_state.get("cdr_errors",pd.DataFrame()).empty:
        st.dataframe(st.session_state.cdr_errors,hide_index=True)

with controls.container(key="camera_panel"):
    st.markdown("### ① เลือกไฟล์")
    st.caption("ใช้ VOICE หรือ DATA อย่างน้อยหนึ่งไฟล์ พร้อมตารางกล้อง")
    cv, cd = st.columns(2)
    camera_voice = cv.file_uploader("ไฟล์ VOICE สำหรับเทียบกล้อง", type=["csv", "txt", "xlsx"], key="camera_voice")
    camera_data = cd.file_uploader("ไฟล์ DATA สำหรับเทียบกล้อง", type=["csv", "txt", "xlsx"], key="camera_data")
    camera_network = st.selectbox("เครือข่าย CDR สำหรับเทียบกล้อง", ["ไม่ระบุ"] + list(NETWORKS), format_func=network_label, key="camera_network")
    st.caption("อัปโหลดข้อมูลกล้องเพื่อเทียบเวลาเหตุการณ์ CDR · ชื่อด่านใช้ข้อความหลัง | · ไฟล์กล้องไม่ถูกบันทึกลงฐานข้อมูล")
    camera_file = st.file_uploader("ไฟล์กล้อง/ด่าน (CSV, XLSX)", type=["csv","txt","xlsx"], key="camera_file")
    st.markdown("### ② กำหนดช่วงเวลา")
    camera_window = st.number_input("ก่อนและหลังเวลากล้อง (นาที)", min_value=1, max_value=240, value=10)
    st.caption("เช่น 10 นาที: กล้องเวลา 10:00 จะเทียบ CDR ตั้งแต่ 09:50 ถึง 10:10")
    camera_run = st.button("จับคู่กล้องกับ CDR", type="primary", disabled=not (camera_file and (camera_voice or camera_data) and (ready or DEMO)))
    if camera_run and camera_file:
        try:
            batches, offset, errors_found = [], 0, []
            for kind, file in [("VOICE", camera_voice), ("DATA", camera_data)]:
                if file:
                    raw = read_table(file.getvalue(), file.name, max_rows=None)
                    events, errors = prepare_cdr(raw, kind=kind, offset=offset, interpret_service=True)
                    offset += len(raw)
                    if not events.empty:
                        batches.append(events)
                    if not errors.empty:
                        errors_found.append(errors)
            if not batches:
                raise ValueError("ไม่พบเหตุการณ์ CDR ที่อ่านได้ในไฟล์ VOICE / DATA")
            events = pd.concat(batches, ignore_index=True).sort_values(["event_at", "event_id"])
            with st.spinner("กำลังค้นหาพิกัด G-Mon สำหรับ CDR ที่เทียบกล้อง…"):
                candidates = demo_rows() if DEMO else db.lookup_pairs(engine, list(zip(events.xci, events.lac)), [] if camera_network == "ไม่ระบุ" else [camera_network])
                if DEMO and camera_network != "ไม่ระบุ":
                    candidates = [r for r in candidates if r["plmn"] == camera_network]
                camera_cdr_rows = match_cdr(events, candidates)
            if errors_found:
                st.warning("มีแถว CDR ที่อ่านไม่ได้ กรุณาตรวจรายละเอียด")
                st.dataframe(pd.concat(errors_found, ignore_index=True), hide_index=True)
            if prepare_camera is None:
                raise ValueError("รุ่นที่ Deploy อยู่ยังไม่รองรับตัวแปลงไฟล์กล้อง กรุณารีเฟรช Deploy")
            camera_raw = read_table(camera_file.getvalue(), camera_file.name, max_rows=None)
            cameras, camera_errors = prepare_camera(camera_raw)
            st.session_state.camera_cdr_rows = camera_cdr_rows
            st.session_state.update(rows=camera_cdr_rows, mode="กล้อง + CDR", elapsed=None, notice=f"วิเคราะห์ CDR สำหรับกล้อง {len(events):,} เหตุการณ์")
            st.session_state.revision += 1
            # Fill camera coordinates from the checkpoint reference table when
            # the event file contains only the checkpoint name.
            try:
                ref = db.checkpoints(engine)
                by_name = {(str(x['checkpoint_name']).strip(), str(x.get('direction') or '').strip()): x for x in ref}
                for i, cam in cameras.iterrows():
                    key = (str(cam.get('checkpoint') or '').strip(), str(cam.get('direction') or '').replace(' กทม.','').replace('จาก ','').strip())
                    hit = by_name.get(key)
                    if hit:
                        cameras.at[i, 'lat'], cameras.at[i, 'lon'] = hit['lat'], hit['lon']
            except Exception:
                pass
            st.session_state.camera_rows = cameras
            st.session_state.camera_generation = st.session_state.get("camera_generation", 0) + 1
            st.session_state.camera_errors = camera_errors
            st.session_state.camera_map_rows = [
                {**r, "observed_at": r.get("camera_time"), "source": "กล้อง", "xci": None, "lac": None, "xnbid": None}
                for r in cameras.to_dict("records") if coordinates(r.get("lat"), r.get("lon"))
            ]
            if cameras.empty:
                st.warning("ไม่พบรายการกล้องที่อ่านได้")
            else:
                st.success(f"อ่านข้อมูลกล้องได้ {len(cameras):,} รายการ")
        except Exception as exc:
            st.error(f"อ่านไฟล์กล้องไม่สำเร็จ: {exc}")
    cameras = st.session_state.get("camera_rows", pd.DataFrame())
    camera_issues = st.session_state.get("camera_errors", pd.DataFrame())
    if active_panel == "camera_panel" and not camera_issues.empty:
        report_area.warning("มีรายการกล้องที่อ่านไม่ได้ กรุณาตรวจแถวและเวลาต่อไปนี้")
        report_area.dataframe(camera_issues, hide_index=True)
    if not cameras.empty and active_panel == "camera_panel":
        report_area.markdown("### ③ ผลวิเคราะห์และไทม์ไลน์")
        report_area.caption("เลือกช่องหน้าแถวเพื่อแสดงกล้องและทุกจุดของเหตุการณ์บนแผนที่")
        with report_area.expander("ดูรายการกล้องที่นำเข้า"):
            st.dataframe(cameras.assign(camera_time=cameras.camera_time.map(display_time)), hide_index=True, width="stretch")
        cdr_events = [r for r in st.session_state.get("camera_cdr_rows", []) if r.get("event_at") and r.get("event_type") != "CAMERA"]
        if cdr_events:
            timeline, unmatched, focus_groups = [], [], []
            matched_cameras = 0
            matched_ids = set()
            for cam in cameras.to_dict("records"):
                nearby = [r for r in cdr_events if abs((r["event_at"] - cam["camera_time"]).total_seconds()) <= camera_window*60]
                # Prefer VOICE at the same window; DATA fills windows with no call.
                voice = [r for r in nearby if str(r.get("event_type", "")).startswith("VOICE")]
                selected = voice or [r for r in nearby if r.get("cdr_kind") in {"DATA", "SMS"}]
                if selected:
                    matched_cameras += 1
                else:
                    unmatched.append({"เวลา": cam["camera_time"], "ทะเบียน": cam.get("plate"), "ด่าน": cam["checkpoint"], "เหตุผล": "ไม่พบ CDR ในช่วงเวลา"})
                if not coordinates(cam.get("lat"), cam.get("lon")):
                    unmatched.append({"เวลา": cam["camera_time"], "ทะเบียน": cam.get("plate"), "ด่าน": cam["checkpoint"], "เหตุผล": "ด่านไม่มีพิกัด"})
                for r in selected:
                    focus_groups.append((cam, r["event_id"]))
                    matched_ids.add(r.get("event_id"))
                    timeline.append({"เวลา": cam["camera_time"], "ทะเบียน": cam.get("plate"), "จังหวัด": cam.get("province"), "กล้อง/ด่าน": cam["checkpoint"],
                                     "ทิศทาง": cam["direction"], "ประเภท": r.get("event_type"),
                                     "Service Type": r.get("service_type"), "กิจกรรม": r.get("activity"),
                                     "เวลา CDR": r["event_at"], "LAC": r.get("lac"), "CELL": r.get("xci"),
                                     "ต่างจากเวลากล้อง (วินาที)": (r["event_at"]-cam["camera_time"]).total_seconds(),
                                     "ละติจูดกล้อง": cam.get("lat"), "ลองจิจูดกล้อง": cam.get("lon"),
                                     "ละติจูด CDR/G-Mon": r.get("lat"), "ลองจิจูด CDR/G-Mon": r.get("lon"),
                                     "แหล่งพิกัด": r.get("source"), "สถานะ": "เวลาใกล้เคียง ไม่ยืนยันการผ่านกล้อง"})
            for event_id in dict.fromkeys(r["event_id"] for r in cdr_events):
                if event_id not in matched_ids:
                    event = next(r for r in cdr_events if r["event_id"] == event_id)
                    unmatched.append({"เวลา": event["event_at"], "ประเภท": event.get("cdr_kind"), "LAC": event.get("lac"), "CELL": event.get("xci"), "เหตุผล": "ไม่ถูกเลือกจับคู่ (นอกช่วงเวลาหรือ VOICE มีลำดับก่อน DATA)"})
            summary_cols = report_area.columns(3)
            summary_cols[0].metric("กล้องที่มี CDR ใกล้เวลา", matched_cameras)
            summary_cols[1].metric("กล้องที่ไม่พบ CDR", len(cameras)-matched_cameras)
            summary_cols[2].metric("รายการกล้องไม่มีพิกัด", sum(not coordinates(r.get("lat"),r.get("lon")) for r in cameras.to_dict("records")))
            if timeline or unmatched:
                # Keep the matched CDR rows available to the map renderer on this run.
                st.session_state.rows = [
                    {**row, "camera_match": row.get("event_id") in matched_ids}
                    for row in cdr_events
                ]
                report_area.success(f"CDR ที่ถูกเลือก {len(matched_ids):,} เหตุการณ์ · รายการจับคู่พิกัด {len(timeline):,} แถว")
                timeline_table = pd.DataFrame(timeline)
                if not timeline_table.empty:
                    timeline_table = timeline_table.sort_values(["เวลา", "เวลา CDR"], kind="stable")
                if not timeline_table.empty:
                    selection = report_area.dataframe(timeline_table, hide_index=True, width="stretch", on_select="rerun", selection_mode="single-row", key=f"camera_timeline_{st.session_state.get('camera_generation', 0)}_{camera_window}_{len(timeline)}")
                    if selection.selection.rows:
                        original_index = timeline_table.index[selection.selection.rows[0]]
                        selected_cam, selected_event = focus_groups[original_index]
                        camera_focus_key = f"{selected_cam.get('camera_id')}_{selected_event}"
                        camera_focus_rows = [r for r in cdr_events if r["event_id"] == selected_event] + [{**selected_cam, "event_type":"CAMERA", "source":"กล้อง"}]
                        report_area.caption("แผนที่แสดงกล้องและทุกพิกัดของเหตุการณ์ที่เลือก · ยกเลิกการเลือกแถวเพื่อแสดงทั้งหมด")
                if unmatched:
                    with report_area.expander("รายการที่ต้องตรวจต่อ"):
                        report_area.dataframe(pd.DataFrame(unmatched), hide_index=True)

                excel_buffer = BytesIO()
                with pd.ExcelWriter(excel_buffer, engine="openpyxl", datetime_format="DD/MM/YYYY HH:MM:SS") as writer:
                    timeline_table.to_excel(writer, sheet_name="ไทม์ไลน์จับคู่", index=False)
                    pd.DataFrame(unmatched, columns=["เวลา", "ทะเบียน", "ด่าน", "ประเภท", "LAC", "CELL", "เหตุผล"]).to_excel(writer, sheet_name="รายการที่ต้องตรวจต่อ", index=False)
                    pd.DataFrame({"เงื่อนไข": ["ช่วงเวลารอบกล้อง (นาที)", "การเลือกประเภท", "ความหมายผลลัพธ์"],
                                  "ค่า": [str(camera_window), "VOICE ก่อน DATA ในช่วงเวลาที่เลือก", "จับคู่ตามเวลา ไม่ยืนยันตำแหน่งโทรศัพท์หรือการผ่านกล้อง"]}).to_excel(writer, sheet_name="เงื่อนไข", index=False)
                    from openpyxl.styles import Font, PatternFill
                    from openpyxl.utils import get_column_letter
                    for sheet in writer.book.worksheets:
                        sheet.freeze_panes = "A2"
                        sheet.auto_filter.ref = sheet.dimensions
                        for cell in sheet[1]:
                            cell.font = Font(bold=True, color="FFFFFF")
                            cell.fill = PatternFill("solid", fgColor="14263D")
                        for column in sheet.columns:
                            sheet.column_dimensions[get_column_letter(column[0].column)].width = min(65, max(23, max(len(str(c.value or "")) for c in column)+2))
                            for cell in column[1:]:
                                if isinstance(cell.value, str):
                                    cell.data_type = "s"
                                elif isinstance(cell.value, datetime):
                                    cell.number_format = "dd/mm/yyyy hh:mm:ss"
                report_area.download_button("📥 ส่งออกไทม์ไลน์ Excel", data=excel_buffer.getvalue(),
                                   file_name="camera_cdr_timeline.xlsx", mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                                   key="export_camera_timeline")
            else:
                report_area.warning("ไม่พบ CDR ในช่วงเวลาที่กำหนดรอบกล้อง")
        else:
            st.info("เลือกไฟล์ VOICE หรือ DATA และไฟล์กล้องในเมนูนี้ แล้วกดจับคู่กล้องกับ CDR")

with controls.container(key="area_panel"):
    with st.form("area"):
        a,b = st.columns(2)
        south = a.number_input("ใต้",-90.,90.,6.9,format="%.5f")
        north = b.number_input("เหนือ",-90.,90.,7.2,format="%.5f")
        west = a.number_input("ตะวันตก",-180.,180.,100.3,format="%.5f")
        east = b.number_input("ตะวันออก",-180.,180.,100.7,format="%.5f")
        area_search = st.form_submit_button("ค้นหาภายในพื้นที่",disabled=not (ready or DEMO))
    if area_search:
        try:
            if south>=north or west>=east:
                raise ValueError("ขอบเขตพื้นที่ไม่ถูกต้อง")
            found = [r for r in demo_rows() if south<=r['lat']<=north and west<=r['lon']<=east] if DEMO else db.search(engine,bounds=dict(south=south,north=north,west=west,east=east))[0]
            st.session_state.update(rows=found,mode="พื้นที่",elapsed=None,notice=f"พบ {len(found):,} รายการในพื้นที่ รวมข้อมูลซ้ำ")
            st.session_state.revision += 1
        except ValueError as exc:
            st.error(str(exc))
        except Exception:
            st.error("ค้นหาพื้นที่ไม่สำเร็จ กรุณาตรวจการเชื่อมต่อ")

rows = st.session_state.rows
if active_panel == "camera_panel":
    rows = camera_focus_rows if camera_focus_rows is not None else rows + st.session_state.get("camera_map_rows", [])
with map_area:
    if st.session_state.notice:
        st.success(st.session_state.notice)
    valid = [r for r in rows if coordinates(r.get("lat"),r.get("lon"))]
    totals = [("พิกัดทั้งหมด", len(valid)), ("G-Mon", sum(r.get('source','G-Mon')=='G-Mon' for r in valid)),
              ("CDR สำรอง", sum(r.get('source')=='CDR สำรอง' for r in valid)), ("ไม่มีพิกัด", len(rows)-len(valid))]
    st.markdown('<div class="result-strip">' + ''.join(f'<div><span>{label}</span><strong>{value:,}</strong></div>' for label,value in totals) + '</div>', unsafe_allow_html=True)
    st.markdown("### 🗺️ แผนที่วิเคราะห์")
    st.caption("📞 VOICE · 📡 DATA · 📷 กล้อง" if active_panel in {"cdr_panel", "camera_panel"} else "🟢 สัญญาณดี · 🟡 ปานกลาง · 🔴 อ่อน · ⚪ ไม่มีเกณฑ์")
    with controls.expander("⚙️ ตัวเลือกแผนที่และคำอธิบาย"):
        st.caption("LTE RSRP: เขียว ≥ −90 · เหลือง ≥ −105 · แดง < −105 dBm · เก็บข้อมูลทุกแถว จุดพิกัดเดียวกันอาจซ้อนกัน")
        playback = st.checkbox("เล่นไทม์ไลน์แบบเน้นกลุ่มจุด",value=True,disabled=st.session_state.mode != "CDR")
        sector = st.checkbox("เปิด Sector จำลอง")
        a,b,c = st.columns(3)
        bearing_text = a.text_input("ทิศจำลอง (0–359°)",placeholder="เช่น 90")
        radius = b.number_input("รัศมีจำลอง (เมตร)",10,20000,1000,100)
        beam = c.number_input("ความกว้าง (องศา)",1,360,90)
        bearing = None
        if bearing_text.strip():
            try:
                bearing = float(bearing_text)
                if not 0<=bearing<360:
                    raise ValueError()
            except ValueError:
                bearing = None
                st.warning("ทิศต้องตั้งแต่ 0 ถึงน้อยกว่า 360 องศา")
        if sector and (bearing is None or len(valid)>500):
            st.info("กรอกทิศจำลอง และค้นหาไม่เกิน 500 จุดเพื่อเปิดรูปพัด")
    timeline_enabled = playback and st.session_state.mode == "CDR" and len(rows) <= 2000
    if playback and st.session_state.mode == "CDR" and len(rows) > 2000:
        st.info("ผลลัพธ์มีจำนวนมาก จึงแสดงจุดทั้งหมดบนแผนที่และปิดตัวควบคุมไทม์ไลน์ชั่วคราวเพื่อความเสถียร")
    map_obj = build_map(rows,sector=sector,bearing=bearing,radius=radius,beam=beam,path=timeline_enabled)
    st_folium(map_obj,height=620,use_container_width=True,key=f"map_{st.session_state.revision}_{camera_focus_key}",returned_objects=[])
    st.caption("จุด G-Mon คือจุดตรวจพบสัญญาณ ไม่ใช่ตำแหน่งเสาจริง · ไทม์ไลน์เน้นทุกจุดของเหตุการณ์โดยไม่เลือกตำแหน่งโทรศัพท์เอง")
    if st.session_state.elapsed is not None:
        st.caption(f"ค้นหาฐานข้อมูล {st.session_state.elapsed:.2f} วินาที")
    if any(r.get("ambiguous") for r in rows):
        st.warning("มีหลาย PLMN ตรงกับ LAC/CELL เดียวกัน แสดงทุกเครือข่าย กรุณาเลือกเครือข่ายเพื่อจำกัดผล")

if rows:
    with st.expander("ตารางผลทั้งหมด · รวมรายการซ้ำ"):
        table = pd.DataFrame(rows)
        for column in ("event_at","observed_at"):
            if column in table:
                table[column] = table[column].map(display_time)
        columns = [c for c in ["event_type","event_id","event_at","source","plmn","xci","lac","lac_display","xnbid","lat","lon","rsrp","observed_at","status"] if c in table]
        st.dataframe(table[columns],hide_index=True,width="stretch")
st.markdown('<div class="footer">วิเคราะห์พิกัด <span>G-Mon เก็บทุกแถว · CDR ใช้เฉพาะเซสชัน</span></div>',unsafe_allow_html=True)
