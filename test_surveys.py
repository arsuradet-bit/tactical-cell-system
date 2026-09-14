import io
import unittest
from datetime import datetime
import pandas as pd
from core import paired_ids, signal_level, prepare_cdr, read_table, match_cdr, display_time
from survey_maps import build_map


class SurveyTests(unittest.TestCase):
    def test_pairs_not_cartesian(self):
        _,_,pairs=paired_ids('117266\n116876\n64918','55653\n55653\n06611')
        self.assertEqual(pairs,[('117266','55653'),('116876','55653'),('64918','6611')])
        with self.assertRaises(ValueError): paired_ids('1\n2','3')
        with self.assertRaises(ValueError): paired_ids('1\n\n2','3\n4\n5')

    def test_signal_boundaries(self):
        for value,level in [(-82,'HIGH'),(-90,'HIGH'),(-90.1,'MID'),(-105,'MID'),(-105.1,'LOW')]:
            self.assertEqual(signal_level(dict(system='4',rsrp=value))[1],level)
        self.assertEqual(signal_level(dict(system='3',rsrp=-82))[0],'#94a3b8')

    def test_excel_underlying_value_and_datetime(self):
        from openpyxl import Workbook
        wb=Workbook(); ws=wb.active
        ws.append(['Start Date','LAC','Cell ID'])
        ws.append([datetime(2025,9,3,9,11,6),'06611',350342102])
        ws['A2'].number_format='yyyy-mm-dd';ws['C2'].number_format='0.0E+00'
        stream=io.BytesIO();wb.save(stream)
        events,errors=prepare_cdr(read_table(stream.getvalue(),'data.xlsx'),kind='DATA')
        self.assertTrue(errors.empty)
        self.assertEqual(events.iloc[0].xci,'350342102')
        self.assertEqual(display_time(events.iloc[0].event_at),'03/09/25 09:11:06')
        self.assertEqual(display_time('03/09/25 09:11:06'),'03/09/25 09:11:06')

    def test_group_playback_preserves_duplicate_surveys(self):
        events,_=prepare_cdr(pd.DataFrame([{'start date':'03/09/25 09:11:06','lac':'6611','cell id':'350342102'},
                                          {'start date':'03/09/25 09:12:06','lac':'99','cell id':'2'}]),kind='DATA')
        a=dict(xci='350342102',lac='6611',plmn='52003',lat=7.,lon=100.,system='4',rsrp=-98)
        candidates=[{**a,'observation_key':'a'},{**a,'observation_key':'b','lat':7.01},{**a,'observation_key':'c'}]
        result=match_cdr(events,candidates)
        self.assertEqual(len(result),4)
        self.assertEqual(result[-1]['source'],'ไม่พบพิกัด')
        markup=build_map(result,path=True).get_root().render()
        self.assertIn('"members": [0, 1, 2]',markup)
        self.assertIn('"members": []',markup)
        self.assertNotIn('L.polyline',markup)


if __name__=='__main__':unittest.main()
